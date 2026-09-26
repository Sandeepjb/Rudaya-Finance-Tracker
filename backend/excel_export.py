"""Excel export of the whole tracker (Transactions, Project P&L, Monthly, Sales Forecast, Quotations, Budgets)."""
import io
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from server import db, get_current_user, _q_totals, _parse_iso

router = APIRouter(prefix="/api/export")
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
HEAD_FILL = PatternFill("solid", fgColor="111111")
HEAD_FONT = Font(bold=True, color="FFFFFF")
INR = '"₹"#,##0.00'
THIN = Side(style="thin", color="DDDDDD")


def _sheet(wb, title, headers, rows, money_cols=(), widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for col in money_cols:
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.number_format = INR
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(12, min(45, len(h) + 6)))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


@router.get("/excel")
async def export_excel(user: dict = Depends(get_current_user), year: Optional[int] = None):
    wb = Workbook()
    wb.remove(wb.active)
    txns = await db.transactions.find({}).sort("date", 1).to_list(50000)

    # Transactions
    _sheet(wb, "Transactions", ["Date", "Type", "Account", "Amount", "Project ID", "Notes", "Source", "Created By"],
           [[(t.get("date") or "")[:10], t.get("type"), t.get("account"), t.get("amount"), t.get("project_id"),
             t.get("notes", ""), t.get("source", ""), t.get("created_by", "")] for t in txns],
           money_cols=(4,), widths={"Notes": 50, "Account": 28, "Project ID": 20})

    # Project P&L
    proj = {}
    for t in txns:
        b = proj.setdefault(t.get("project_id") or "(none)", {"Revenue": 0, "Cost": 0, "Expense": 0})
        b[t["type"]] = b.get(t["type"], 0) + t["amount"]
    prows = [[p, b["Revenue"], b["Cost"], b["Expense"], b["Revenue"] - b["Cost"], b["Revenue"] - b["Cost"] - b["Expense"]]
             for p, b in sorted(proj.items(), key=lambda kv: -(kv[1]["Revenue"] + kv[1]["Cost"] + kv[1]["Expense"]))]
    ws = _sheet(wb, "Project P&L", ["Project ID", "Revenue", "Cost", "Expense", "Gross Profit", "Net Profit"], prows,
                money_cols=(2, 3, 4, 5, 6), widths={"Project ID": 24})
    n = len(prows) + 2
    ws.append(["TOTAL"] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{n - 1})" for c in range(2, 7)])
    for c in ws[n]:
        c.font = Font(bold=True)
        if c.column > 1:
            c.number_format = INR

    # Monthly Revenue vs Expense (all years present)
    monthly = {}
    for t in txns:
        dt = _parse_iso(t["date"])
        if not dt:
            continue
        b = monthly.setdefault((dt.year, dt.month), {"Revenue": 0, "Cost": 0, "Expense": 0})
        b[t["type"]] = b.get(t["type"], 0) + t["amount"]
    mrows = [[y, MONTHS[m - 1], b["Revenue"], b["Cost"], b["Expense"], b["Revenue"] - b["Cost"] - b["Expense"]]
             for (y, m), b in sorted(monthly.items())]
    _sheet(wb, "Monthly", ["Year", "Month", "Revenue", "Cost", "Expense", "Net"], mrows, money_cols=(3, 4, 5, 6))

    # Sales Forecast
    sf = await db.sales_forecast.find({} if not year else {"year": year}).sort([("year", 1), ("month", 1)]).to_list(20000)
    _sheet(wb, "Sales Forecast", ["Year", "Month", "Type", "Project ID", "Amount", "Notes"],
           [[d["year"], MONTHS[d["month"] - 1], d.get("type", "Revenue"), d.get("project_id", ""), d["amount"], d.get("notes", "")]
            for d in sf], money_cols=(5,), widths={"Notes": 40, "Project ID": 20})

    # Quotations + lines
    qs = await db.quotations.find({}).sort("quote_date", 1).to_list(5000)
    _sheet(wb, "Quotations", ["Quotation No", "Client", "Project ID", "Quote Date", "Expected", "Status", "Revenue", "Cost",
                              "Expense", "Net", "Notes"],
           [[q.get("quotation_number"), q.get("client_name", ""), q.get("project_id", ""), q.get("quote_date", ""),
             f'{q.get("expected_year")}-{(q.get("expected_month") or 0):02d}', q.get("status", ""),
             *(lambda t: [t["Revenue"], t["Cost"], t["Expense"], t["net"]])(_q_totals(q.get("lines") or [])), q.get("notes", "")]
            for q in qs], money_cols=(7, 8, 9, 10), widths={"Client": 28, "Notes": 40})
    _sheet(wb, "Quotation Lines", ["Quotation No", "Type", "Description", "Amount", "Notes"],
           [[q.get("quotation_number"), ln.get("type"), ln.get("description", ""), ln.get("amount"), ln.get("notes", "")]
            for q in qs for ln in (q.get("lines") or [])], money_cols=(4,), widths={"Description": 45, "Notes": 30})

    # Budgets vs actual
    budgets = await db.budgets.find({} if not year else {"year": year}).sort([("year", 1), ("account", 1)]).to_list(5000)
    actual_by = {}
    for t in txns:
        dt = _parse_iso(t["date"])
        if dt and t["type"] == "Expense":
            actual_by[(dt.year, t["account"])] = actual_by.get((dt.year, t["account"]), 0) + t["amount"]
    _sheet(wb, "Expense Budget", ["Year", "Account", "Annual Budget", "Actual Expense", "Variance"],
           [[b.get("year"), b.get("account"), b.get("amount", b.get("annual_budget", 0)),
             actual_by.get((b.get("year"), b.get("account")), 0),
             b.get("amount", b.get("annual_budget", 0)) - actual_by.get((b.get("year"), b.get("account")), 0)] for b in budgets],
           money_cols=(3, 4, 5), widths={"Account": 28})

    # Master data
    accounts = await db.accounts.find({}, {"_id": 0}).to_list(1000)
    pids = await db.project_ids.find({}, {"_id": 0}).to_list(1000)
    _sheet(wb, "Master Data", ["Kind", "Value", "Description"],
           [["Account", a["name"], ""] for a in accounts] + [["Project ID", p["code"], p.get("description", "")] for p in pids])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="Rudaya-Finance-Tracker-{stamp}.xlsx"'})
