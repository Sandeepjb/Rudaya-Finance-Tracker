"""Tests: statement CSV import, email alert parsing templates, bulk approve, Excel export."""
import io
import os
import re
import uuid
from pathlib import Path

import pytest
import requests
from dotenv import dotenv_values

BASE_URL = (os.environ.get("REACT_APP_BACKEND_URL")
            or dotenv_values("/app/frontend/.env")["REACT_APP_BACKEND_URL"]).rstrip("/")
API = f"{BASE_URL}/api"
BT = f"{API}/bank-transactions"
ENV = dotenv_values("/app/backend/.env")
HDR = {"X-Ingest-Key": ENV["BANK_INGEST_API_KEY"]}
W = "".join(chr(65 + int(c)) if c.isdigit() else c for c in os.environ.get("PYTEST_XDIST_WORKER", "gw0").upper())  # per-worker tag so parallel module teardowns never collide


@pytest.fixture(scope="module", autouse=True)
def cleanup():
    yield
    import asyncio
    from motor.motor_asyncio import AsyncIOMotorClient
    from bson import ObjectId
    cl = AsyncIOMotorClient(ENV["MONGO_URL"])
    d = cl[ENV["DB_NAME"]]

    async def run():
        fin_ids = [ObjectId(x["finance_transaction_id"]) async for x in d.bank_transactions.find(
            {"$or": [{"bank_name": {"$regex": "^QAEXT" + W}}, {"narration": {"$regex": "QAEXT" + W}}]}, {"finance_transaction_id": 1}) if x.get("finance_transaction_id")]
        if fin_ids:
            await d.transactions.delete_many({"_id": {"$in": fin_ids}})
        await d.transactions.delete_many({"source": "bank_transaction", "notes": {"$regex": "^QAEXT" + W}})
        await d.bank_transactions.delete_many({"bank_name": {"$regex": "^QAEXT" + W}})
        await d.bank_transactions.delete_many({"narration": {"$regex": "QAEXT" + W}})
        await d.bank_parse_templates.delete_many({"bank_name": {"$regex": "^QAEXT" + W}})
        await d.bank_mapping_rules.delete_many({"pattern": {"$regex": "^QAEXT" + W}})
    asyncio.get_event_loop().run_until_complete(run())


@pytest.fixture(scope="module")
def admin():
    c = Path("/app/memory/test_credentials.md").read_text()
    e = re.search(r"(?im)^\s*(?:[-*]\s*)?(?:\*\*)?email(?:\*\*)?\s*:\s*`?([^`\s]+)", c).group(1)
    pw = re.search(r"(?im)^\s*(?:[-*]\s*)?(?:\*\*)?password(?:\*\*)?\s*:\s*`?([^`\s]+)", c).group(1)
    s = requests.Session()
    assert s.post(f"{API}/auth/login", json={"email": e, "password": pw}).status_code == 200
    return s


@pytest.fixture(scope="module")
def meta(admin):
    m = admin.get(f"{API}/meta").json()
    return {"account": m["accounts"][0]["name"], "project": m["project_ids"][0]["code"]}


def _tag():
    return "QAEXT" + W + "".join(chr(65 + int(c, 16) % 26) for c in uuid.uuid4().hex[:5])


ICICI_CSV = ("Account Statement\nCustomer Name: QA\n\n"
             "Txn Date,Value Date,Transaction Remarks,Chq/Ref No,Withdrawal Amount (INR),Deposit Amount (INR),Balance\n"
             "12-05-2026,12-05-2026,{t} SUBSCRIPTION 998,{r}A,\"4,499.13\",,100000\n"
             "13-05-2026,13-05-2026,NEFT CR {t} CLIENT,{r}B,,\"70,000.00\",170000\n"
             "bad-date,13-05-2026,{t} BROKEN,{r}C,10,,170000\n"
             "14-05-2026,14-05-2026,{t} ZERO,{r}D,,,170000\n")
HDFC_CSV = ("Date,Narration,Chq./Ref.No.,Value Dt,Withdrawal Amt.,Deposit Amt.,Closing Balance\n"
            "25/09/26,UPI-{t} LEI CODE REGISTRATION,{r}X,25/09/26,15222.00,0.00,50000.00\n")
GENERIC_CSV = ("Date,Description,Amount,Dr/Cr,Reference\n"
               "2026-06-01,{t} GENERIC PAY,2500,DR,{r}G\n"
               "2026-06-02,{t} GENERIC RECEIPT,9000,CR,{r}H\n")


def _upload(admin, path, csv_text, **fields):
    return admin.post(f"{BT}/statement/{path}", files={"file": ("st.csv", io.BytesIO(csv_text.encode()), "text/csv")}, data=fields)


class TestStatementCsv:
    @pytest.fixture(autouse=True)
    def _lock(self, txn_state_lock):
        yield

    def test_icici_preview_detects_bank_and_columns(self, admin):
        t, r = _tag(), _tag()
        p = _upload(admin, "preview", ICICI_CSV.format(t=t, r=r)).json()
        assert p["bank_name"] == "ICICI Bank"
        assert p["mapping"]["date"] == "Txn Date" and p["mapping"]["narration"] == "Transaction Remarks"
        assert p["mapping"]["debit"].startswith("Withdrawal") and p["mapping"]["credit"].startswith("Deposit")
        assert p["total"] == 4 and p["valid"] == 2 and p["invalid"] == 2
        rows = {x["row"]: x for x in p["rows"]}
        assert rows[1]["txn"]["direction"] == "debit" and rows[1]["txn"]["amount"] == 4499.13
        assert rows[2]["txn"]["direction"] == "credit" and rows[2]["txn"]["amount"] == 70000
        assert "unparseable date" in rows[3]["errors"] and any("amount" in e for e in rows[4]["errors"])

    def test_hdfc_preview(self, admin):
        t, r = _tag(), _tag()
        p = _upload(admin, "preview", HDFC_CSV.format(t=t, r=r)).json()
        assert p["bank_name"] == "HDFC Bank" and p["valid"] == 1
        assert p["rows"][0]["txn"]["transaction_date"] == "2026-09-25" and p["rows"][0]["txn"]["direction"] == "debit"

    def test_generic_amount_drcr_and_manual_mapping(self, admin):
        t, r = _tag(), _tag()
        p = _upload(admin, "preview", GENERIC_CSV.format(t=t, r=r), bank_name="QAEXT" + W + " Bank").json()
        assert p["bank_name"] == "QAEXT" + W + " Bank" and p["valid"] == 2
        dirs = [x["txn"]["direction"] for x in p["rows"]]
        assert dirs == ["debit", "credit"]
        # manual remap: point narration at Reference column
        p2 = _upload(admin, "preview", GENERIC_CSV.format(t=t, r=r), bank_name="QAEXT" + W + " Bank",
                     mapping='{"date":"Date","narration":"Reference","amount":"Amount","direction":"Dr/Cr"}').json()
        assert p2["rows"][0]["txn"]["narration"] == f"{r}G"

    def test_import_runs_pipeline_and_skips_reingest(self, admin):
        t, r = _tag(), _tag()
        csv_text = ICICI_CSV.format(t=t, r=r)
        res = _upload(admin, "import", csv_text, bank_name="QAEXT" + W + " ICICI").json()
        assert res["ingested"] == 2 and res["invalid"] == 2 and res["pending"] == 2
        for x in res["results"]:
            assert x["source"] == "csv" and x["status"] == "pending" and x["suggestion"] is not None
        again = _upload(admin, "preview", csv_text, bank_name="QAEXT" + W + " ICICI").json()
        assert again["already_ingested"] == 2
        res2 = _upload(admin, "import", csv_text, bank_name="QAEXT" + W + " ICICI").json()
        assert res2["ingested"] == 0 and res2["skipped_already_ingested"] == 2
        h = admin.get(f"{BT}/ingestion-history").json()
        assert any(x["channel"] == "csv_statement" for x in h)

    def test_bad_file(self, admin):
        assert _upload(admin, "preview", "hello,world\n1,2\n").status_code == 422
        assert _upload(admin, "preview", "   ").status_code == 422

    def test_non_admin_blocked(self):
        s = requests.Session()
        assert s.post(f"{BT}/statement/preview", files={"file": ("a.csv", b"x", "text/csv")}).status_code == 401


class TestEmailTemplates:
    @pytest.fixture(autouse=True)
    def _lock(self, txn_state_lock):
        yield

    def test_default_templates_seeded(self, admin):
        tp = admin.get(f"{BT}/templates").json()
        banks = {t["bank_name"] for t in tp}
        assert {"ICICI Bank", "HDFC Bank", "Saraswat Bank"} <= banks
        for t in tp:
            if t["is_default"]:
                r = admin.post(f"{BT}/templates/test", json={"raw_text": t["sample"], "template_id": t["id"]}).json()
                assert r["errors"] == [], (t["name"], r)
                assert r["fields"]["amount"] > 0 and r["fields"]["direction"] in ("debit", "credit")

    def test_icici_alert_parse(self, admin):
        text = "Dear Customer, Your Account XX2345 has been debited with INR 4,499.13 on 12-May-26. Info: MICROSOFT SUBSCRIPTION E080105RXG. The Available Balance is INR 1,20,000.00."
        r = admin.post(f"{BT}/templates/test", json={"raw_text": text, "bank_name": "ICICI Bank"}).json()
        f = r["fields"]
        assert f["amount"] == 4499.13 and f["direction"] == "debit" and f["transaction_date"] == "2026-05-12"
        assert f["narration"].startswith("MICROSOFT SUBSCRIPTION") and f["bank_account"] == "2345"

    def test_hdfc_credit_alert_parse(self, admin):
        text = "Rs.70000.00 credited to a/c **6789 on 13-05-26 by NEFT from RUDAYA CLIENT LTD (UTR HDFCN52026051312345). Avl bal Rs.1,00,000.00"
        r = admin.post(f"{BT}/templates/test", json={"raw_text": text, "bank_name": "HDFC Bank"}).json()
        f = r["fields"]
        assert f["direction"] == "credit" and f["amount"] == 70000 and f["bank_reference"] == "HDFCN52026051312345"

    def test_email_ingest_endpoint(self, admin):
        t = _tag()
        mid = "pa-" + uuid.uuid4().hex
        text = f"Rs.15222.00 debited from a/c **6789 on 25-09-26 to {t} LEI CODE REGISTRATION (UPI Ref No {t}5268). Not you? Call 18002586161."
        r = requests.post(f"{BT}/ingest/email", headers=HDR, json={"bank_name": "HDFC Bank", "raw_text": text, "source_message_id": mid})
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["template"] == "HDFC UPI/NEFT alert"
        assert d["parsed"]["narration"] == f"{t} LEI CODE REGISTRATION"
        assert d["transaction"]["status"] == "pending" and d["transaction"]["source"] == "power_automate"
        assert d["transaction"]["bank_account_masked"].endswith("6789")
        # replay with same message id → duplicate
        r2 = requests.post(f"{BT}/ingest/email", headers=HDR, json={"bank_name": "HDFC Bank", "raw_text": text, "source_message_id": mid}).json()
        assert r2["transaction"]["status"] == "duplicate"

    def test_email_ingest_unparseable_logs_failure(self, admin):
        r = requests.post(f"{BT}/ingest/email", headers=HDR, json={"bank_name": "ICICI Bank", "raw_text": "Hello this is not a bank alert at all QAEXT", "source_message_id": "x" + uuid.uuid4().hex})
        assert r.status_code == 422
        assert "errors" in r.json()["detail"]
        h = admin.get(f"{BT}/ingestion-history").json()
        assert any(x["channel"] == "email_alert" and x.get("failed") for x in h)

    def test_email_ingest_requires_key(self):
        assert requests.post(f"{BT}/ingest/email", json={"raw_text": "Rs.1 debited on 01-01-26 to X"}).status_code == 401

    def test_template_crud_and_validation(self, admin):
        bank = _tag()
        body = {"bank_name": bank, "name": "custom", "priority": 1, "sample": "Amt INR 99.50 was paid on 01-02-2026 for QAEXT THING. Ref: ABC12345",
                "patterns": {"amount": r"INR\s*([\d,\.]+)", "debit": r"(paid)", "credit": r"(received)", "date": r"on\s+(\d{2}-\d{2}-\d{4})",
                             "narration": r"for\s+([^\.]+)", "reference": r"Ref:\s*(\w+)", "account": ""}}
        c = admin.post(f"{BT}/templates", json=body)
        assert c.status_code == 200, c.text
        tid = c.json()["id"]
        r = admin.post(f"{BT}/templates/test", json={"raw_text": body["sample"], "template_id": tid}).json()
        assert r["errors"] == [] and r["fields"]["narration"] == "QAEXT THING" and r["fields"]["bank_reference"] == "ABC12345"
        bad = dict(body, patterns=dict(body["patterns"], amount="([unclosed"))
        assert admin.put(f"{BT}/templates/{tid}", json=bad).status_code == 422
        u = admin.put(f"{BT}/templates/{tid}", json=dict(body, enabled=False)).json()
        assert u["enabled"] is False
        assert admin.delete(f"{BT}/templates/{tid}").status_code == 200
        assert admin.delete(f"{BT}/templates/{tid}").status_code == 404

    def test_guide(self, admin):
        g = admin.get(f"{BT}/guide").json()
        assert g["header_name"] == "X-Ingest-Key" and g["email_endpoint"].endswith("/api/bank-transactions/ingest/email")
        assert len(g["steps"]) >= 5 and ENV["BANK_INGEST_API_KEY"] not in str(g)


class TestBulkApprove:
    @pytest.fixture(autouse=True)
    def _lock(self, txn_state_lock):
        yield

    def test_bulk_approve_mixed(self, admin, meta):
        t = _tag()
        ids = []
        for i in range(3):
            r = requests.post(f"{BT}/ingest", headers=HDR, json={"transactions": [{
                "bank_name": "QAEXT" + W + " Bulk", "transaction_date": "2026-05-01", "direction": "debit", "amount": 100 + i,
                "narration": f"{t} BULK {i}", "bank_reference": f"{t}{i}"}]}).json()["results"][0]
            ids.append(r["id"])
        for i in ids[:2]:
            admin.put(f"{BT}/{i}", json={"type": "Expense", "account": meta["account"], "project_id": meta["project"], "amount": 100 + ids.index(i), "date": "2026-05-01", "notes": "QAEXT" + W + f" bulk {i}"})
        # third: force incomplete suggestion by editing to invalid master value is blocked; instead reject it so it's not pending
        admin.post(f"{BT}/{ids[2]}/reject", json={"reason": "qa"})
        res = admin.post(f"{BT}/bulk-approve", json={"ids": ids + ["000000000000000000000000"]}).json()
        assert res["requested"] == 4 and res["approved"] == 2 and res["failed"] == 2, res
        ok = [x for x in res["results"] if x["ok"]]
        assert all(x["finance_transaction_id"] for x in ok)
        fails = [x for x in res["results"] if not x["ok"]]
        assert any("rejected" in x["error"].lower() for x in fails) and any("not found" in x["error"].lower() for x in fails)
        for i in ids[:2]:
            assert admin.get(f"{BT}/{i}").json()["status"] == "approved"
        # re-run is idempotent
        res2 = admin.post(f"{BT}/bulk-approve", json={"ids": ids[:2]}).json()
        assert res2["approved"] == 0

    def test_bulk_requires_admin(self):
        assert requests.post(f"{BT}/bulk-approve", json={"ids": ["x"]}).status_code == 401


class TestExcelExport:
    def test_export_workbook(self, admin):
        r = admin.get(f"{API}/export/excel")
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers["content-type"]
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        names = wb.sheetnames
        for s in ("Transactions", "Project P&L", "Monthly", "Sales Forecast", "Quotations", "Quotation Lines", "Expense Budget", "Master Data"):
            assert s in names, names
        ws = wb["Transactions"]
        assert [c.value for c in ws[1]][:5] == ["Date", "Type", "Account", "Amount", "Project ID"]
        assert ws.max_row > 1

    def test_export_requires_auth(self):
        assert requests.get(f"{API}/export/excel").status_code == 401
